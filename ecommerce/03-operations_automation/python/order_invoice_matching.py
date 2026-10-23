"""Extract order numbers from invoice descriptions and build a mapping workbook."""
import re
from pathlib import Path
import openpyxl

INPUT_FILE = Path("invoices.xlsx")
OUTPUT_FILE = Path("orders_and_invoices.xlsx")
SHEET_NAME = "invoices"
ORDER_PATTERN = re.compile(r"order\D*(\d{9,})", re.IGNORECASE)


def build_mapping(input_file: Path = INPUT_FILE, output_file: Path = OUTPUT_FILE) -> int:
    if not input_file.exists():
        raise FileNotFoundError(f"Input workbook not found: {input_file.resolve()}")

    workbook = openpyxl.load_workbook(input_file)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"Worksheet {SHEET_NAME!r} not found; available: {workbook.sheetnames}")
    invoice_sheet = workbook[SHEET_NAME]
    if invoice_sheet.max_column < 4:
        raise ValueError("Invoice sheet must contain invoice number in column B and description in column D")

    result_sheet = (
        workbook["Orders_and_Invoices"]
        if "Orders_and_Invoices" in workbook.sheetnames
        else workbook.create_sheet("Orders_and_Invoices")
    )
    if result_sheet.max_row:
        result_sheet.delete_rows(1, result_sheet.max_row)
    result_sheet.cell(1, 1, "Order Number")
    result_sheet.cell(1, 2, "Invoice Number")

    result_row = 2
    seen: set[tuple[str, str]] = set()
    for row in invoice_sheet.iter_rows(min_row=2, values_only=True):
        invoice_number = str(row[1]).strip() if row[1] is not None else ""
        description = str(row[3]).strip() if row[3] is not None else ""
        for order_number in ORDER_PATTERN.findall(description):
            pair = (order_number, invoice_number)
            if pair in seen:
                continue
            seen.add(pair)
            result_sheet.cell(result_row, 1, order_number)
            result_sheet.cell(result_row, 2, invoice_number)
            result_row += 1

    workbook.save(output_file)
    return result_row - 2


if __name__ == "__main__":
    print(f"Created {build_mapping()} mapping row(s) in {OUTPUT_FILE.resolve()}")
